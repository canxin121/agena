"""Generate small local reference fixtures for the ignored document adapter test.

Run with the same Python environment that provides MarkItDown's Office extras.
"""
import pathlib
import sys
import zipfile

target = pathlib.Path(sys.argv[1])
target.mkdir(parents=True, exist_ok=True)
with zipfile.ZipFile(target / "test.docx", "w") as document:
    document.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    document.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    document.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Fixture needle 中文</w:t></w:r></w:p></w:body></w:document>')

from openpyxl import Workbook
workbook = Workbook()
workbook.active.append(["Name", "Value"])
workbook.active.append(["Fixture needle 中文", 42])
workbook.save(target / "test.xlsx")

from pptx import Presentation
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[0])
slide.shapes.title.text = "Fixture needle 中文"
presentation.save(target / "test.pptx")

content = b"BT /F1 14 Tf 50 700 Td (Fixture needle PDF) Tj ET"
objects = [
    b"<< /Type /Catalog /Pages 2 0 R >>",
    b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
    b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
    b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
]
pdf = bytearray(b"%PDF-1.4\n")
offsets = []
for index, obj in enumerate(objects, 1):
    offsets.append(len(pdf))
    pdf.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
xref = len(pdf)
pdf.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
for offset in offsets:
    pdf.extend(f"{offset:010d} 00000 n \n".encode())
pdf.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
(target / "test.pdf").write_bytes(pdf)
print(target)
